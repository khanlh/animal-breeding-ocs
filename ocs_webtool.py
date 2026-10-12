#!/usr/bin/env python3
"""
Pig breeding workflow web tool (runs locally):  BLUP results -> selection index -> optiSel

Start it from the folder where you normally run the scripts:

    python ocs_webtool.py

and the browser opens http://127.0.0.1:8765

The tool only calls your two scripts (nothing is re-implemented here):
    step 1: si_cal.py        (python)
    step 2: optiSel_pig.R    (Rscript)
Keep this file in the same folder as the two scripts, or point to them with
--scripts-dir / --si-script / --ocs-script.

Relative paths typed in the page are resolved from the folder in which this tool
was started, exactly like in the terminal (e.g. ../data_260929/blup_LL_T2613).
Only the Python standard library is needed by the tool itself (si_cal.py needs
pandas, openpyxl and matplotlib). The server listens on 127.0.0.1 only.
"""

import argparse
import ast
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import uuid
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlparse

WORKDIR = Path.cwd()
DEFAULT_TRAITS = ["D105", "BF105", "Loin105", "Live", "BWt", "Wt21"]
DEFAULT_DIRECTIONS = {"D105": -1, "BF105": -1, "Loin105": 1, "Live": 1, "BWt": 1, "Wt21": 1}
CFG = {"python": sys.executable, "rscript": None, "si_script": None, "ocs_script": None,
       "default_traits": list(DEFAULT_TRAITS), "default_directions": dict(DEFAULT_DIRECTIONS),
       "max_traits": 12, "supports_traits": False, "supports_weights": False}
JOBS = {}
ALLOWED = []          # folders whose files may be served
LOCK = threading.Lock()
TRAIT_NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.\-]*$")


# ---------------------------------------------------------------- helpers
def resolve(p):
    p = Path(os.path.expanduser(str(p)))
    if not p.is_absolute():
        p = WORKDIR / p
    return p.resolve()


def allow(path):
    path = Path(path).resolve()
    with LOCK:
        if path not in ALLOWED:
            ALLOWED.append(path)


def is_allowed(path):
    path = Path(path).resolve()
    with LOCK:
        roots = list(ALLOWED)
    for r in roots:
        try:
            if os.path.commonpath([str(r), str(path)]) == str(r):
                return True
        except ValueError:
            continue
    return False


def in_dir(base, name):
    """File name relative to base, unless it is already absolute."""
    if re.match(r"^(/|~|[A-Za-z]:)", name):
        return Path(os.path.expanduser(name))
    return Path(base) / name


def find_script(scripts_dir, names):
    for n in names:
        p = Path(scripts_dir) / n
        if p.is_file():
            return str(p)
    return None


def read_script_config(script):
    """Read the trait defaults and the supported options from the SI script."""
    traits, directions, max_traits = list(DEFAULT_TRAITS), dict(DEFAULT_DIRECTIONS), 12
    supports_traits = supports_weights = False
    try:
        src = Path(script).read_text(encoding="utf-8", errors="replace")
        supports_traits = "--traits" in src
        supports_weights = "--weights" in src and "--directions" in src
        found = {}
        for node in ast.parse(src).body:
            if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
                found[node.targets[0].id] = node.value
        for key, target in (("DEFAULT_TRAITS", "traits"), ("TRAITS", "traits")):
            if key in found and target == "traits":
                val = ast.literal_eval(found[key])
                if isinstance(val, list) and val and all(isinstance(x, str) for x in val):
                    traits = val
                    break
        for key in ("DEFAULT_DIRECTION", "DIRECTION"):
            if key in found:
                val = ast.literal_eval(found[key])
                if isinstance(val, dict):
                    directions = {str(k): (-1 if int(v) < 0 else 1) for k, v in val.items()}
                    break
        if "MAX_TRAITS" in found:
            max_traits = int(ast.literal_eval(found["MAX_TRAITS"]))
    except Exception:
        pass
    directions = {t: directions[t] for t in directions}
    return traits, directions, max_traits, supports_traits, supports_weights


def find_rscript(explicit):
    if explicit:
        return explicit
    found = shutil.which("Rscript")
    if found:
        return found
    for c in ("/usr/local/bin/Rscript", "/opt/homebrew/bin/Rscript",
              "/Library/Frameworks/R.framework/Resources/bin/Rscript",
              "/usr/bin/Rscript"):
        if Path(c).is_file():
            return c
    return None


def fmt_cell(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.5g}" if abs(v) < 1000 else f"{v:.1f}"
    s = str(v)
    try:
        f = float(s)
        if re.search(r"[.eE]", s):
            return f"{f:.5g}" if abs(f) < 1000 else f"{f:.1f}"
    except ValueError:
        pass
    return s


def read_csv_table(path, title, max_rows=None):
    with open(path, newline="", encoding="utf-8", errors="replace") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        return None
    head, body = rows[0], rows[1:]
    total = len(body)
    if max_rows:
        body = body[:max_rows]
    return {"title": title, "columns": head,
            "rows": [[fmt_cell(c) for c in r] for r in body], "total": total}


def read_xlsx_table(path, sheet, title, max_rows=15, cols=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        return None
    it = wb[sheet].iter_rows(values_only=True)
    head = [str(h) if h is not None else "" for h in next(it, [])]
    idx = list(range(len(head)))
    if cols:
        idx = [head.index(c) for c in cols if c in head]
    body = []
    for r in it:
        body.append([fmt_cell(r[i]) if i < len(r) else "" for i in idx])
        if max_rows and len(body) >= max_rows:
            break
    wb.close()
    return {"title": title, "columns": [head[i] for i in idx], "rows": body, "total": None}


def xlsx_trait_columns(path):
    """Trait names that have both EBV_<name> and ACC_<name> columns (file order)."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        header = next(wb[wb.sheetnames[0]].iter_rows(values_only=True, max_row=1), [])
        wb.close()
    except Exception:
        return None
    cols = [str(h) for h in header if h is not None]
    acc = {c[4:] for c in cols if c.startswith("ACC_")}
    names = [c[4:] for c in cols if c.startswith("EBV_") and c[4:] in acc]
    return names or None


def matrix_dims(text):
    """{effect: matrix size} for every 'Genetic variance(s) for effect N' block."""
    dims = {}
    for m in re.finditer(r"Genetic\s+variance\(s\)\s+for\s+effect\s+(\d+)", text, flags=re.I):
        size = None
        for line in text[m.end():m.end() + 30000].splitlines():
            if not line.strip():
                continue
            try:
                vals = [float(t.replace("D", "E").replace("d", "e")) for t in line.split()]
            except ValueError:
                if size is not None:
                    break
                continue
            size = len(vals)
            break
        dims[int(m.group(1))] = size
    return dims


# ---------------------------------------------------------------- scan
def scan_dir(d):
    p = resolve(d)
    if not p.is_dir():
        return {"ok": False, "error": f"Folder not found: {p}"}
    allow(p)
    files = sorted(f for f in os.listdir(p) if (p / f).is_file() and not f.startswith("."))

    logs = [f for f in files if f.lower().endswith(".log")]
    effects, dims = {}, {}
    for f in logs:
        try:
            txt = (p / f).read_text(errors="replace")
        except OSError:
            continue
        found = matrix_dims(txt)
        if found:
            effects[f] = sorted(found)
            dims[f] = {str(k): v for k, v in found.items()}

    renadd = []
    for f in files:
        m = re.match(r"renadd0*(\d+)\.ped$", f, flags=re.I)
        if m:
            renadd.append({"file": f, "effect": int(m.group(1))})

    dats = [f for f in files if f.lower().endswith(".dat") and not f.lower().startswith("renf90")]
    xlsx = [f for f in files if f.lower().endswith(".xlsx") and not f.startswith("~$")]
    peds = [f for f in files if f.lower().endswith(".ped") and not f.lower().startswith("renadd")]

    ebv_columns = {}
    for f in xlsx:
        names = xlsx_trait_columns(p / f)
        if names:
            ebv_columns[f] = names

    def first(cands, preferred=()):
        for pr in preferred:
            if pr in cands:
                return pr
        return cands[0] if cands else ""

    log = first(logs, ("blupf90.log",))
    ebv_cands = [f for f in xlsx if "_SI" not in f.upper() and "TREND" not in f.upper()]
    si_cands = [f for f in xlsx if "_SI" in f.upper()]
    animal_hint = renadd[0]["effect"] if len(renadd) == 1 else None
    if animal_hint is not None and effects.get(log) and animal_hint not in effects[log]:
        animal_hint = None

    return {
        "ok": True, "path": str(p), "files": files,
        "logs": logs, "effects": effects, "dims": dims, "renadd": renadd, "animal_hint": animal_hint,
        "dats": dats, "xlsx": xlsx, "peds": peds, "ebv_columns": ebv_columns,
        "suggest": {
            "log": log,
            "ebv": first(ebv_cands, ("EBV_ACC.xlsx",)),
            "dat": dats[0] if len(dats) == 1 else "",
            "ped": first(peds),
            "si": first(si_cands, ("EBV_ACC_SI.xlsx",)),
        },
    }


# ---------------------------------------------------------------- jobs
def bad(msg):
    raise ValueError(msg)


def get_str(params, key, default="", required=False, label=None):
    v = str(params.get(key, "") or "").strip()
    if not v:
        if required:
            bad(f"Missing: {label or key}")
        return default
    return v


def get_num(params, key, kind=int, required=False, label=None, lo=None):
    v = str(params.get(key, "") or "").strip()
    if not v:
        if required:
            bad(f"Missing: {label or key}")
        return None
    try:
        n = kind(float(v)) if kind is int else float(v)
        if kind is int and float(v) != n:
            raise ValueError
    except ValueError:
        bad(f"{label or key} must be {'an integer' if kind is int else 'a number'}: '{v}'")
    if lo is not None and n < lo:
        bad(f"{label or key} must be >= {lo}")
    return n


def parse_traits(p):
    """Validate the trait table sent by the page -> (names, weights, directions)."""
    rows = p.get("traits") or []
    mx = CFG["max_traits"]
    if not isinstance(rows, list) or not 1 <= len(rows) <= mx:
        bad(f"Between 1 and {mx} traits are required.")
    names, weights, dirs = [], {}, {}
    for i, r in enumerate(rows, 1):
        name = str(r.get("name", "")).strip()
        if not TRAIT_NAME_RE.match(name):
            bad(f"Trait {i}: invalid name '{name}'. Use letters, digits, '_', '.', '-' "
                "(no spaces, commas or '='), starting with a letter or digit.")
        if name in names:
            bad(f"Trait name '{name}' is used more than once.")
        raw = str(r.get("weight", "")).strip()
        try:
            w = 1.0 if raw == "" else float(raw)
        except ValueError:
            bad(f"Weight of {name} must be a number: '{raw}'")
        if w != w or w in (float("inf"), float("-inf")) or w < 0:
            bad(f"Weight of {name} must be >= 0.")
        d = str(r.get("direction", "")).strip()
        if d in ("-1", "-"):
            dirs[name] = -1
        elif d in ("1", "+1", "+"):
            dirs[name] = 1
        else:
            bad(f"Select the selection direction of {name} (lower or higher is better).")
        names.append(name)
        weights[name] = w
    if sum(weights.values()) <= 0:
        bad("At least one weight must be greater than 0.")
    return names, weights, dirs


def build_si(p):
    if not CFG["si_script"]:
        bad("si_cal.py (or si_all.py) not found. Use --scripts-dir or --si-script.")
    d = resolve(get_str(p, "dir", required=True, label="BLUP results folder"))
    if not d.is_dir():
        bad(f"Folder not found: {d}")
    log = get_str(p, "log", "blupf90.log")
    ebv = get_str(p, "ebv", "EBV_ACC.xlsx")
    dat = get_str(p, "dat", required=True, label=".dat file")
    outx = get_str(p, "out_xlsx", "EBV_ACC_SI.xlsx")
    outp = get_str(p, "out_pdf", "EBV_trend.pdf")
    eff = get_num(p, "effect", int, required=True, label="Effect", lo=1)
    for label, name in (("Log file", log), ("EBV file", ebv), (".dat file", dat)):
        f = in_dir(d, name)
        if not f.is_file():
            bad(f"{label} does not exist: {f}")
    names, weights, dirs = parse_traits(p)

    argv = [CFG["python"], "-u", CFG["si_script"], "--base-dir", str(d), "--ebv-file", ebv,
            "--log-file", log, "--dat-file", dat, "--genetic-effect", str(eff),
            "--out-xlsx", outx, "--out-pdf", outp]
    for key, flag, kind, lo in (("ref_min_year", "--ref-min-year", int, 1900), ("min_acc", "--min-acc", float, 0),
                                ("dat_id_col", "--dat-id-col", int, 1), ("dat_year_col", "--dat-year-col", int, 1),
                                ("min_n", "--min-n", int, 1)):
        n = get_num(p, key, kind, label=flag, lo=lo)
        if n is not None:
            argv += [flag, str(n)]
    sd = get_str(p, "sd_mode", "group")
    if sd not in ("group", "genetic"):
        bad("sd-mode must be group or genetic")
    argv += ["--sd-mode", sd]

    acc_traits = [t for t in re.split(r"[,;\s]+", get_str(p, "acc_traits")) if t]
    if acc_traits:
        unknown = [t for t in acc_traits if t not in names]
        if unknown:
            bad(f"ACC filter traits not in the trait table: {', '.join(unknown)}")
        if get_num(p, "min_acc", float, lo=0) is not None:
            argv += ["--acc-traits"] + acc_traits

    defaults_same = (names == CFG["default_traits"] and weights == {t: 1.0 for t in names}
                     and all(dirs[t] == CFG["default_directions"].get(t) for t in names))
    if CFG["supports_traits"]:
        argv += ["--traits", ",".join(names),
                 "--weights", ",".join(f"{t}={weights[t]:g}" for t in names),
                 "--directions", ",".join(f"{t}={dirs[t]:+d}" for t in names)]
    elif names == CFG["default_traits"] and CFG["supports_weights"]:
        argv += ["--weights", ",".join(f"{t}={weights[t]:g}" for t in names),
                 "--directions", ",".join(f"{t}={dirs[t]:+d}" for t in names)]
    elif not defaults_same:
        bad("This si_cal.py is an old version and does not support custom trait names. Replace it with "
            "the new version, or keep the default traits (weights and directions can still be changed "
            "if the script supports --weights).")
    allow(d)
    return argv, {"dir": str(d), "out_xlsx": str(in_dir(d, outx)), "out_pdf": str(in_dir(d, outp)),
                  "traits": names}


def build_ocs(p):
    if not CFG["ocs_script"]:
        bad("optiSel_pig.R (or ocs_pig.R) not found. Use --scripts-dir or --ocs-script.")
    if not CFG["rscript"]:
        bad("Rscript not found. Use --rscript /path/to/Rscript.")
    d = resolve(get_str(p, "dir", required=True, label="input-dir"))
    if not d.is_dir():
        bad(f"Folder not found: {d}")
    ped = get_str(p, "ped", required=True, label="ped-file")
    si = get_str(p, "si", required=True, label="si-file")
    sires = get_num(p, "sires", int, required=True, label="Number of sires", lo=1)
    dams = get_num(p, "dams", int, required=True, label="Number of dams", lo=1)
    year = get_str(p, "birth_year")
    cand = get_str(p, "cand")
    if not year and not cand:
        bad("Give a birth year (or 'all') or a cand-file.")
    for label, name in (("ped-file", ped), ("si-file", si), ("cand-file", cand)):
        if name and not in_dir(d, name).is_file():
            bad(f"{label} does not exist: {in_dir(d, name)}")
    argv = [CFG["rscript"], CFG["ocs_script"], "--input-dir", str(d), "--ped-file", ped,
            "--si-file", si, "--sires", str(sires), "--dams", str(dams)]
    if cand:
        argv += ["--cand-file", cand]
    elif year.lower() == "all":
        argv += ["--birth-year", "all"]
    else:
        argv += ["--birth-year", str(get_num(p, "birth_year", int, label="birth-year", lo=1900))]
    for key, flag, kind, lo in (("min_matings", "--min-matings", int, 0), ("max_matings", "--max-matings", int, 1),
                                ("ne", "--ne", float, 1), ("pop_size", "--pop-size", int, 1),
                                ("alternatives", "--alternatives", int, 1)):
        n = get_num(p, key, kind, label=flag, lo=lo)
        if n is not None:
            argv += [flag, str(n)]
    for key, flag in (("si_column", "--si-column"), ("sheet", "--sheet"), ("out_root", "--out-root"),
                      ("out_name", "--out-name"), ("tag", "--tag")):
        v = get_str(p, key)
        if v:
            argv += [flag, v]
    allow(d)
    return argv, {"dir": str(d)}


def collect_results(job):
    meta, step, lines = job["meta"], job["step"], job["lines"]
    outputs, tables = [], []
    try:
        if step == "si":
            for key, label in (("out_xlsx", "SI ranking (Excel)"), ("out_pdf", "EBV trend charts (PDF)")):
                f = Path(meta[key])
                if f.is_file():
                    outputs.append({"name": f.name, "label": label, "path": str(f)})
            if Path(meta["out_xlsx"]).is_file():
                cols = ["Rank", "ID", "Year", "SI_Raw", "SI_100", "Percentile"] + [f"EBV_{t}" for t in meta["traits"]]
                t = read_xlsx_table(meta["out_xlsx"], "Ranked_Animals", "Top 15 animals by SI", 15, cols)
                if t:
                    tables.append(t)
                t = read_xlsx_table(meta["out_xlsx"], "Index_Summary", "Index summary", None)
                if t:
                    tables.append(t)
        else:
            outdir = None
            for ln in lines:
                m = re.match(r"\s*Output folder:\s*(.+?)\s*$", ln)
                if m:
                    outdir = Path(m.group(1))
            if outdir and outdir.is_dir():
                allow(outdir)
                meta["outdir"] = str(outdir)
                for f in sorted(os.listdir(outdir)):
                    outputs.append({"name": f, "label": "", "path": str(outdir / f)})
                specs = (("6_QC_summary.csv", "QC summary", None), ("4_sire_mating_counts.csv", "Sires used", None),
                         ("3_mating_plan.csv", "Mating plan (first 15 rows)", 15))
                for fn, title, mx in specs:
                    f = outdir / fn
                    if f.is_file():
                        t = read_csv_table(f, title, mx)
                        if t:
                            tables.append(t)
    except Exception as exc:  # preview problems must not hide the job result
        lines.append(f"[webtool] Could not read the results for the preview: {exc}")
    job["outputs"], job["tables"] = outputs, tables


def run_job(job):
    env = dict(os.environ, PYTHONUNBUFFERED="1")
    try:
        proc = subprocess.Popen(job["argv"], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                                bufsize=1, cwd=str(WORKDIR), env=env, errors="replace")
        job["proc"] = proc
        for line in proc.stdout:
            job["lines"].append(line.rstrip("\n"))
        proc.wait()
        job["rc"] = proc.returncode
    except Exception as exc:
        job["lines"].append(f"[webtool] Could not start the command: {exc}")
        job["rc"] = -1
    if job.get("stopped"):
        job["status"] = "stopped"
    else:
        if job["rc"] == 0:
            collect_results(job)
        job["status"] = "done" if job["rc"] == 0 else "error"
    job["finished"] = time.time()


def start_job(step, params):
    argv, meta = (build_si if step == "si" else build_ocs)(params)
    job = {"id": uuid.uuid4().hex[:10], "step": step, "argv": argv, "meta": meta, "lines": [],
           "status": "running", "rc": None, "outputs": [], "tables": [], "started": time.time()}
    job["lines"].append("$ " + " ".join(f'"{a}"' if " " in a else a for a in argv))
    JOBS[job["id"]] = job
    threading.Thread(target=run_job, args=(job,), daemon=True).start()
    return job["id"]


# ---------------------------------------------------------------- HTTP
class Handler(BaseHTTPRequestHandler):
    server_version = "OCSWebtool/2.0"

    def log_message(self, fmt, *args):
        pass

    def _host_ok(self):
        return self.headers.get("Host", "").split(":")[0] in ("127.0.0.1", "localhost")

    def _send(self, body, ctype, code=200, extra=None):
        if isinstance(body, str):
            body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, obj, code=200):
        self._send(json.dumps(obj, ensure_ascii=False), "application/json; charset=utf-8", code)

    def do_GET(self):
        if not self._host_ok():
            return self._send("Forbidden", "text/plain", 403)
        u = urlparse(self.path)
        q = parse_qs(u.query)
        if u.path == "/":
            return self._send(HTML, "text/html; charset=utf-8")
        if u.path == "/api/env":
            return self._json({
                "python": CFG["python"], "rscript": CFG["rscript"], "workdir": str(WORKDIR),
                "si_script": CFG["si_script"], "ocs_script": CFG["ocs_script"],
                "default_traits": CFG["default_traits"], "default_directions": CFG["default_directions"],
                "max_traits": CFG["max_traits"], "supports_traits": CFG["supports_traits"],
                "supports_weights": CFG["supports_weights"]})
        if u.path == "/api/job":
            job = JOBS.get((q.get("id") or [""])[0])
            if not job:
                return self._json({"error": "Unknown job"}, 404)
            off = int((q.get("offset") or ["0"])[0])
            lines = job["lines"][off:]
            out = {"status": job["status"], "lines": lines, "next": off + len(lines)}
            if job["status"] in ("done", "error", "stopped"):
                out.update({"rc": job["rc"], "outputs": job["outputs"], "tables": job["tables"],
                            "meta": job["meta"], "seconds": round(job.get("finished", time.time()) - job["started"], 1)})
            return self._json(out)
        if u.path == "/file":
            path = Path((q.get("path") or [""])[0])
            if not path.is_absolute() or not path.is_file() or not is_allowed(path):
                return self._send("Not allowed or not found", "text/plain; charset=utf-8", 404)
            ext = path.suffix.lower()
            ctype = {".pdf": "application/pdf", ".csv": "text/plain; charset=utf-8",
                     ".txt": "text/plain; charset=utf-8", ".xlsx":
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}.get(ext, "application/octet-stream")
            inline = ext in (".pdf", ".txt") or (q.get("inline") or ["0"])[0] == "1"
            disp = f"{'inline' if inline else 'attachment'}; filename*=UTF-8''{quote(path.name)}"
            return self._send(path.read_bytes(), ctype, extra={"Content-Disposition": disp})
        return self._send("Not found", "text/plain", 404)

    def do_POST(self):
        if not self._host_ok():
            return self._send("Forbidden", "text/plain", 403)
        if "application/json" not in self.headers.get("Content-Type", ""):
            return self._json({"error": "Content-Type must be application/json"}, 415)
        try:
            n = int(self.headers.get("Content-Length", "0"))
            data = json.loads(self.rfile.read(n) or b"{}")
        except (ValueError, json.JSONDecodeError):
            return self._json({"error": "Invalid JSON"}, 400)
        path = urlparse(self.path).path
        try:
            if path == "/api/scan":
                return self._json(scan_dir(str(data.get("dir", ""))))
            if path == "/api/run":
                step = data.get("step")
                if step not in ("si", "ocs"):
                    return self._json({"error": "step must be si or ocs"}, 400)
                return self._json({"job": start_job(step, data.get("params") or {})})
            if path == "/api/stop":
                job = JOBS.get(str(data.get("id", "")))
                if job and job["status"] == "running" and job.get("proc"):
                    job["stopped"] = True
                    job["proc"].terminate()
                return self._json({"ok": True})
        except ValueError as exc:
            return self._json({"error": str(exc)}, 400)
        return self._json({"error": "Not found"}, 404)


# ---------------------------------------------------------------- page
HTML = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Pig OCS workflow</title>
<style>
:root{--bg:#f6f7f9;--card:#fff;--ink:#1c2230;--mute:#667085;--line:#e3e6ec;--accent:#1f6feb;--ok:#1a7f4b;--err:#c62828;--warn:#9a6700;--code:#0f1623;--codeink:#d6deeb;--chip:#eef2ff}
@media (prefers-color-scheme:dark){:root:not([data-theme=light]){--bg:#10141c;--card:#171c26;--ink:#e6e9f0;--mute:#98a2b3;--line:#262d3a;--accent:#5b9dff;--ok:#4cc38a;--err:#ff7b72;--warn:#e3b341;--code:#0b0f17;--chip:#1f2a44}}
:root[data-theme=dark]{--bg:#10141c;--card:#171c26;--ink:#e6e9f0;--mute:#98a2b3;--line:#262d3a;--accent:#5b9dff;--ok:#4cc38a;--err:#ff7b72;--warn:#e3b341;--code:#0b0f17;--chip:#1f2a44}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Helvetica,Arial,sans-serif}
.wrap{max-width:980px;margin:0 auto;padding:24px 16px 64px}
h1{font-size:22px;margin:0 0 4px}
.sub{color:var(--mute);margin:0 0 16px}
.chips{display:flex;flex-wrap:wrap;gap:8px;margin-bottom:18px}
.chip{background:var(--chip);border:1px solid var(--line);border-radius:999px;padding:3px 10px;font-size:12.5px;color:var(--ink)}
.chip.bad{border-color:var(--err);color:var(--err)}
.chip.good::before{content:"● ";color:var(--ok)}
.chip.bad::before{content:"● ";color:var(--err)}
.steps{display:flex;align-items:center;gap:10px;margin:0 0 18px;flex-wrap:wrap}
.step{display:flex;align-items:center;gap:8px;padding:6px 12px;border:1px solid var(--line);border-radius:10px;background:var(--card);font-weight:600;font-size:14px}
.step .dot{width:10px;height:10px;border-radius:50%;background:var(--line)}
.step.running .dot{background:var(--accent);animation:pulse 1s infinite}
.step.done .dot{background:var(--ok)}
.step.error .dot{background:var(--err)}
.arrow{color:var(--mute)}
@keyframes pulse{50%{opacity:.3}}
.card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:18px 18px 16px;margin-bottom:18px}
.card h2{font-size:17px;margin:0 0 2px}
.card .desc{color:var(--mute);margin:0 0 14px;font-size:13.5px}
.grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px 16px}
.grid .full{grid-column:1/-1}
@media (max-width:700px){.grid{grid-template-columns:1fr}}
label{display:block;font-size:13px;font-weight:600;margin-bottom:3px}
label small{font-weight:400;color:var(--mute);font-family:ui-monospace,Menlo,monospace;margin-left:4px}
input,select{width:100%;padding:8px 10px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--ink);font:inherit}
input:focus,select:focus{outline:2px solid var(--accent);outline-offset:-1px}
.hint{font-size:12.5px;color:var(--mute);margin-top:3px}
.row{display:flex;gap:8px}
.row input{flex:1}
button{font:inherit;cursor:pointer;border-radius:8px;border:1px solid var(--line);background:var(--card);color:var(--ink);padding:8px 14px}
button.primary{background:var(--accent);border-color:var(--accent);color:#fff;font-weight:600}
button.small{padding:4px 10px;font-size:13px}
button:disabled{opacity:.5;cursor:not-allowed}
button.ef{padding:3px 10px;font-size:12.5px;border-radius:999px;background:var(--chip)}
button.ef.sel{border-color:var(--accent);outline:1px solid var(--accent)}
details{margin-top:12px;border-top:1px dashed var(--line);padding-top:10px}
summary{cursor:pointer;color:var(--mute);font-size:13.5px}
.actions{display:flex;gap:10px;margin-top:16px;align-items:center;flex-wrap:wrap}
.status{font-size:13.5px;color:var(--mute)}
.status.ok{color:var(--ok)}.status.err{color:var(--err)}
.msg{margin-top:12px;padding:9px 12px;border-radius:8px;font-size:13.5px;border:1px solid var(--line);background:var(--bg);display:none}
.msg.err{display:block;border-color:var(--err);color:var(--err)}
.msg.info{display:block}
pre.log{background:var(--code);color:var(--codeink);border-radius:10px;padding:12px;margin:14px 0 0;max-height:300px;overflow:auto;font:12.5px/1.45 ui-monospace,Menlo,monospace;display:none;white-space:pre-wrap;word-break:break-all}
pre.log .e{color:#ff7b72}pre.log .w{color:#e3b341}pre.log .c{color:#7ee787}
.res{margin-top:16px;display:none}
.res h3{font-size:14px;margin:16px 0 6px}
.tw{overflow-x:auto;border:1px solid var(--line);border-radius:10px}
table{border-collapse:collapse;width:100%;font-size:12.5px}
th,td{padding:5px 9px;border-bottom:1px solid var(--line);text-align:left;white-space:nowrap}
th{background:var(--bg);position:sticky;top:0}
tr:last-child td{border-bottom:0}
.files{display:flex;flex-wrap:wrap;gap:8px}
.files a{display:inline-block;padding:5px 11px;border:1px solid var(--line);border-radius:8px;text-decoration:none;color:var(--accent);font-size:13px;background:var(--bg)}
.outprev{font-size:12.5px;color:var(--mute);margin-top:4px;font-family:ui-monospace,Menlo,monospace;word-break:break-all}
.theme{position:fixed;top:12px;right:14px}
.traits{margin-top:16px}
.traits .head{display:flex;justify-content:space-between;align-items:center;gap:8px;flex-wrap:wrap}
#tt td{padding:4px 8px}
#tt input{padding:4px 8px}
#tt input.t-name{width:150px}
#tt input.t-w{width:90px}
#tt select{width:auto;padding:4px 8px}
#tt td.pct,#tt td.idx{color:var(--mute);font-variant-numeric:tabular-nums}
.dimcheck div{margin-top:4px;font-size:12.5px}
.dimcheck .ok{color:var(--ok)}.dimcheck .err{color:var(--err)}
</style>
</head>
<body>
<button class="theme" id="theme" title="Toggle light/dark theme">◐</button>
<div class="wrap">
  <h1>Animal breeding workflow</h1>
  <p class="sub">BLUP results &rarr; selection index &rarr; optiSel (optimum contributions + mating list)</p>
  <div class="chips" id="env"></div>

  <div class="steps">
    <div class="step" id="st1"><span class="dot"></span>1 · Selection index</div><span class="arrow">→</span>
    <div class="step" id="st2"><span class="dot"></span>2 · optiSel</div>
  </div>

  <!-- ============ STEP 1 ============ -->
  <div class="card" id="card1">
    <h2>1 · Compute the selection index from BLUP results</h2>
    <p class="desc">Runs <code>si_cal.py</code>: reads the genetic covariance matrix from the <code>.log</code> file and the EBVs from Excel, computes the index and draws EBV trends by birth year.</p>
    <div class="grid">
      <div class="full">
        <label>BLUP results folder <small>--base-dir</small></label>
        <div class="row"><input id="s1_dir" data-save placeholder="../data_260929/blup_LL_T2613"><button id="scan">Scan folder</button></div>
        <div class="hint" id="scanmsg">Type a path (relative to the folder where the tool was started, or absolute) and click “Scan folder” to detect the files, effects and traits.</div>
      </div>
      <div class="full">
        <label>Effect of interest (animal effect) <small>--genetic-effect</small></label>
        <input id="s1_effect" data-save type="number" min="1" placeholder="e.g. 7" style="max-width:160px">
        <div class="hint" id="efhint">Folder not scanned yet.</div>
        <div id="efchips" style="display:flex;gap:6px;margin-top:6px;flex-wrap:wrap"></div>
      </div>
      <div><label>Log file <small>--log-file</small></label><input id="s1_log" data-save list="dl_log" value="blupf90.log"></div>
      <div><label>EBV file (Excel) <small>--ebv-file</small></label><input id="s1_ebv" data-save list="dl_xlsx" value="EBV_ACC.xlsx"></div>
      <div><label>Data file (.dat) <small>--dat-file</small></label><input id="s1_dat" data-save list="dl_dat" placeholder="datLL_T2613.dat"></div>
      <div><label>Output SI file <small>--out-xlsx</small></label><input id="s1_outx" data-save value="EBV_ACC_SI.xlsx"></div>
    </div>

    <div class="traits">
      <div class="head">
        <label style="margin:0">Traits, selection direction and weights <small>--traits / --directions / --weights</small></label>
        <span class="status" id="tcount"></span>
      </div>
      <div class="hint">Rows must follow the order of trait 1…N in the BLUP model (the rows of the covariance matrix in the log). Each name must match the Excel columns <code>EBV_&lt;name&gt;</code> and <code>ACC_&lt;name&gt;</code>. Weights are per 1 genetic SD (so traits with different units are comparable); 0 removes a trait from the index.</div>
      <div class="tw" style="margin-top:6px"><table id="tt"><thead><tr><th>#</th><th>Trait name</th><th>Selection direction</th><th>Weight</th><th>Share</th></tr></thead><tbody></tbody></table></div>
      <div class="actions" style="margin-top:8px">
        <button type="button" class="small" id="tadd">+ Add trait</button>
        <button type="button" class="small" id="tdel">− Remove last</button>
        <button type="button" class="small" id="tdetect">Detect from Excel / log</button>
        <button type="button" class="small" id="wreset">Reset weights to 1</button>
        <span class="hint" id="wwarn"></span>
      </div>
      <div class="dimcheck" id="dimcheck"></div>
    </div>

    <details>
      <summary>Advanced options</summary>
      <div class="grid" style="margin-top:10px">
        <div><label>Reference group for SI_100: born in or after <small>--ref-min-year</small></label><input id="s1_refyear" data-save type="number" placeholder="empty = whole herd"></div>
        <div><label>SI_100 scale <small>--sd-mode</small></label>
          <select id="s1_sd" data-save><option value="group">group — SD of the reference group</option><option value="genetic">genetic — genetic SD of the index</option></select></div>
        <div><label>Minimum ACC <small>--min-acc</small></label><input id="s1_minacc" data-save type="number" step="0.01" min="0" max="1" placeholder="empty = no filter"></div>
        <div><label>Traits the ACC filter applies to <small>--acc-traits</small></label><input id="s1_acctraits" data-save placeholder="empty = all traits (comma-separated names)"></div>
        <div><label>Minimum animals per year (charts) <small>--min-n</small></label><input id="s1_minn" data-save type="number" min="1" placeholder="10"></div>
        <div><label>ID column in .dat <small>--dat-id-col</small></label><input id="s1_idcol" data-save type="number" min="1" placeholder="1"></div>
        <div><label>Birth-year column in .dat <small>--dat-year-col</small></label><input id="s1_ycol" data-save type="number" min="1" placeholder="4"></div>
        <div><label>Output PDF file <small>--out-pdf</small></label><input id="s1_outp" data-save value="EBV_trend.pdf"></div>
      </div>
    </details>
    <div class="actions">
      <button class="primary" id="run1">Run step 1</button><button id="stop1" disabled>Stop</button>
      <span class="status" id="stat1"></span>
    </div>
    <div class="msg" id="msg1"></div>
    <pre class="log" id="log1"></pre>
    <div class="res" id="res1"></div>
  </div>

  <!-- ============ STEP 2 ============ -->
  <div class="card" id="card2">
    <h2>2 · optiSel: select sires and dams and build the mating list</h2>
    <p class="desc">Runs <code>optiSel_pig.R</code>: shortlists the top-SI sires and dams, optimises contributions under a kinship ceiling derived from Ne, then pairs animals with the lowest parental kinship.</p>
    <div class="grid">
      <div class="full"><label>Input folder <small>--input-dir</small></label><input id="s2_dir" data-save placeholder="../data_260929/blup_LL_T2613"></div>
      <div><label>Pedigree file <small>--ped-file</small></label><input id="s2_ped" data-save list="dl_ped" placeholder="PedLL.ped"></div>
      <div><label>SI file (output of step 1) <small>--si-file</small></label><input id="s2_si" data-save list="dl_si" value="EBV_ACC_SI.xlsx"></div>
      <div><label>Number of sires to use <small>--sires</small></label><input id="s2_sires" data-save type="number" min="1" value="40"></div>
      <div><label>Number of dams to mate <small>--dams</small></label><input id="s2_dams" data-save type="number" min="1" value="400"></div>
      <div><label>Birth year of the candidate cohort <small>--birth-year</small></label><input id="s2_year" data-save placeholder="2026 or all" value="2026"></div>
      <div><label>List of available animals (optional) <small>--cand-file</small></label><input id="s2_cand" data-save placeholder="alive.csv — if given, the birth year is ignored"></div>
      <div class="full"><div class="outprev" id="outprev"></div></div>
    </div>
    <details>
      <summary>Advanced options</summary>
      <div class="grid" style="margin-top:10px">
        <div><label>Minimum matings per sire <small>--min-matings</small></label><input id="s2_minm" data-save type="number" min="0" placeholder="1"></div>
        <div><label>Maximum matings per sire <small>--max-matings</small></label><input id="s2_maxm" data-save type="number" min="1" placeholder="30"></div>
        <div><label>Target effective size Ne <small>--ne</small></label><input id="s2_ne" data-save type="number" min="1" placeholder="50"></div>
        <div><label>SI column in the Excel file <small>--si-column</small></label><input id="s2_sicol" data-save placeholder="SI_100"></div>
        <div><label>Alternative sires listed per dam <small>--alternatives</small></label><input id="s2_alt" data-save type="number" min="1" placeholder="2"></div>
        <div><label>Population size used by optiSel <small>--pop-size</small></label><input id="s2_pop" data-save type="number" min="1" placeholder="1000"></div>
        <div><label>Excel sheet <small>--sheet</small></label><input id="s2_sheet" data-save placeholder="default: first sheet"></div>
        <div><label>Line/breed label (output folder name) <small>--tag</small></label><input id="s2_tag" data-save placeholder="taken from the folder name, e.g. LL"></div>
        <div><label>Folder for the results <small>--out-root</small></label><input id="s2_root" data-save placeholder="default: folder where the tool was started"></div>
        <div><label>Name of the results folder <small>--out-name</small></label><input id="s2_name" data-save placeholder="default: <parent folder>_<line>_OCS"></div>
      </div>
    </details>
    <div class="actions">
      <button class="primary" id="run2">Run step 2</button><button id="stop2" disabled>Stop</button>
      <span class="status" id="stat2"></span>
    </div>
    <div class="msg" id="msg2"></div>
    <pre class="log" id="log2"></pre>
    <div class="res" id="res2"></div>
  </div>
</div>

<datalist id="dl_log"></datalist><datalist id="dl_xlsx"></datalist><datalist id="dl_dat"></datalist>
<datalist id="dl_ped"></datalist><datalist id="dl_si"></datalist>

<script>
const $ = id => document.getElementById(id);
const esc = s => String(s).replace(/[&<>"]/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
let ENV = {}, SCAN = null, JOB = {1:null, 2:null}, TIMER = {1:null, 2:null};
let TR = [];                 // trait table: [{name, dir, w}]
let MEM = {};                // remembered {name: {dir, w}} across projects

async function api(url, body){
  const opt = body ? {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify(body)} : {};
  const r = await fetch(url, opt);
  let j = {};
  try { j = await r.json(); } catch(e) {}
  if(!r.ok) throw new Error(j.error || r.statusText);
  return j;
}

/* ---------- saved form values ---------- */
function saveForm(){
  try{
    const o = {};
    document.querySelectorAll('[data-save]').forEach(e => o[e.id] = e.value);
    localStorage.setItem('ocs_webtool_v2', JSON.stringify(o));
  }catch(e){}
}
function loadForm(){
  try{
    const o = JSON.parse(localStorage.getItem('ocs_webtool_v2') || '{}');
    for(const k in o){ const e = $(k); if(e && e.hasAttribute('data-save')) e.value = o[k]; }
  }catch(e){}
}
function saveTraits(){
  try{
    TR.forEach(t => { if(t.name.trim()) MEM[t.name.trim()] = {dir: t.dir, w: t.w}; });
    localStorage.setItem('ocs_traits_v2', JSON.stringify(TR));
    localStorage.setItem('ocs_traits_mem', JSON.stringify(MEM));
  }catch(e){}
}
function loadTraits(){
  try{ MEM = JSON.parse(localStorage.getItem('ocs_traits_mem') || '{}'); }catch(e){ MEM = {}; }
  let saved = null;
  try{ saved = JSON.parse(localStorage.getItem('ocs_traits_v2') || 'null'); }catch(e){}
  if(Array.isArray(saved) && saved.length && saved.length <= ENV.max_traits &&
     saved.every(t => t && typeof t.name === 'string')){
    TR = saved.map(t => ({name: t.name, dir: String(t.dir || ''), w: String(t.w === undefined ? '1' : t.w)}));
  }else{
    TR = ENV.default_traits.map(n => ({name: n, dir: dirFor(n), w: '1'}));
  }
}
function dirFor(name){
  if(MEM[name] && MEM[name].dir) return String(MEM[name].dir);
  const d = ENV.default_directions[name];
  return d === undefined ? '' : String(d < 0 ? -1 : 1);
}

/* ---------- environment ---------- */
function chip(text, ok){ return `<span class="chip ${ok?'good':'bad'}">${esc(text)}</span>`; }
async function loadEnv(){
  ENV = await api('/api/env');
  const base = p => p ? p.split('/').pop() : null;
  $('env').innerHTML =
    chip('Python', true) +
    chip(ENV.rscript ? 'Rscript: ' + ENV.rscript : 'Rscript not found', !!ENV.rscript) +
    chip(ENV.si_script ? 'Step 1: ' + base(ENV.si_script) : 'si_cal.py missing', !!ENV.si_script) +
    chip(ENV.ocs_script ? 'Step 2: ' + base(ENV.ocs_script) : 'optiSel_pig.R missing', !!ENV.ocs_script) +
    `<span class="chip">Working folder: ${esc(ENV.workdir)}</span>`;
  if(!ENV.supports_traits) $('env').innerHTML += chip('si_cal.py is an old version: custom traits disabled', false);
  if(!$('s2_root').placeholder.includes('/')) $('s2_root').placeholder = 'default: ' + ENV.workdir;
}

/* ---------- traits table ---------- */
function renderTraits(){
  const tb = document.querySelector('#tt tbody'); tb.innerHTML = '';
  TR.forEach((t, i) => {
    const tr = document.createElement('tr'); tr.dataset.i = i;
    tr.innerHTML = `<td class="idx">${i+1}</td>
      <td><input class="t-name" maxlength="40" placeholder="e.g. ADG" value="${esc(t.name)}"></td>
      <td><select class="t-dir"><option value="">Select…</option><option value="-1">Lower is better (−)</option><option value="1">Higher is better (+)</option></select></td>
      <td><input class="t-w" type="number" min="0" step="0.1" value="${esc(t.w)}"></td><td class="pct"></td>`;
    tb.appendChild(tr);
    tr.querySelector('.t-dir').value = t.dir;
  });
  const editable = ENV.supports_traits;
  tb.querySelectorAll('input,select').forEach(e => {
    e.disabled = !editable && !(ENV.supports_weights && !e.classList.contains('t-name'));
    e.addEventListener('input', traitEdited);
  });
  $('tadd').disabled = !editable || TR.length >= ENV.max_traits;
  $('tdel').disabled = !editable || TR.length <= 1;
  $('tdetect').disabled = !editable;
  $('wreset').disabled = !(ENV.supports_weights || ENV.supports_traits);
  $('tcount').textContent = `${TR.length} / ${ENV.max_traits} traits`;
  updateShares(); dimCheck();
}
function syncFromDom(){
  document.querySelectorAll('#tt tbody tr').forEach(tr => {
    const i = +tr.dataset.i;
    TR[i].name = tr.querySelector('.t-name').value;
    TR[i].dir = tr.querySelector('.t-dir').value;
    TR[i].w = tr.querySelector('.t-w').value;
  });
}
function traitEdited(ev){
  syncFromDom();
  if(ev && ev.target.classList.contains('t-name')){
    const tr = ev.target.closest('tr'), i = +tr.dataset.i, n = TR[i].name.trim();
    if(!TR[i].dir && n){ const d = dirFor(n); if(d){ TR[i].dir = d; tr.querySelector('.t-dir').value = d; } }
  }
  updateShares(); dimCheck(); saveTraits();
}
function updateShares(){
  const rows = [...document.querySelectorAll('#tt tbody tr')];
  const vals = rows.map(tr => { const v = parseFloat(tr.querySelector('.t-w').value); return isFinite(v) && v >= 0 ? v : 0; });
  const sum = vals.reduce((a,b) => a+b, 0);
  rows.forEach((tr, i) => { tr.querySelector('.pct').textContent = sum > 0 ? (100*vals[i]/sum).toFixed(1) + '%' : '—'; });
  const w = sum > 0 ? '' : '<span style="color:var(--err)">At least one weight must be greater than 0.</span>';
  if($('wwarn').innerHTML !== w) $('wwarn').innerHTML = w;
}
function setTraitCount(n){
  syncFromDom();
  while(TR.length > n) TR.pop();
  while(TR.length < n) TR.push({name: '', dir: '', w: '1'});
  renderTraits(); saveTraits();
}
function setTraitNames(names){
  syncFromDom();
  const old = {}; TR.forEach(t => { if(t.name.trim()) old[t.name.trim()] = t; });
  TR = names.map(n => { const o = old[n] || MEM[n] || {}; 
    return {name: n, dir: String(o.dir || dirFor(n) || ''), w: String(o.w === undefined ? '1' : o.w)}; });
  renderTraits(); saveTraits();
}
function resetWeights(){ syncFromDom(); TR.forEach(t => t.w = '1'); renderTraits(); saveTraits(); }

function currentDim(){
  if(!SCAN || !SCAN.dims) return null;
  const log = $('s1_log').value.trim(), eff = String($('s1_effect').value || '').trim();
  const d = SCAN.dims[log]; return d && d[eff] ? d[eff] : null;
}
function detectTraits(){
  if(!SCAN) return 'Scan the folder first.';
  const cols = SCAN.ebv_columns && SCAN.ebv_columns[$('s1_ebv').value.trim()];
  const n = currentDim();
  if(cols && cols.length){
    setTraitNames(cols);
    return `Set ${cols.length} trait names from the EBV_/ACC_ columns of ${$('s1_ebv').value.trim()}.`;
  }
  if(n){
    syncFromDom(); const names = TR.map(t => t.name.trim());
    while(names.length > n) names.pop();
    while(names.length < n) names.push('Trait' + (names.length+1));
    setTraitNames(names);
    return `No EBV_/ACC_ columns found; set ${n} rows from the log matrix size (please type the trait names).`;
  }
  return 'Nothing to detect: no EBV_/ACC_ columns and no matrix in the selected log/effect.';
}
function dimCheck(){
  const box = $('dimcheck');
  if(!SCAN){ if(box.innerHTML){ box.innerHTML = ''; box.dataset.html = ''; } return; }
  const parts = [];
  const n = currentDim(), eff = $('s1_effect').value;
  if(n){
    if(n === TR.length) parts.push(`<div class="ok">Log matrix of effect ${esc(eff)} is ${n}×${n}: matches the ${TR.length} traits in the table.</div>`);
    else parts.push(`<div class="err">Log matrix of effect ${esc(eff)} is ${n}×${n} (${n} traits) but the table has ${TR.length} rows. ` +
      (ENV.supports_traits ? `<button type="button" class="small" data-fix="${n}">Use ${n} rows</button>` : '') + '</div>');
  }
  const file = $('s1_ebv').value.trim(), cols = SCAN.ebv_columns && SCAN.ebv_columns[file];
  if(cols){
    const miss = TR.map(t => t.name.trim()).filter(x => x && !cols.includes(x));
    if(miss.length) parts.push(`<div class="err">No EBV_/ACC_ columns in ${esc(file)} for: ${miss.map(esc).join(', ')}.</div>`);
    else if(TR.every(t => t.name.trim())) parts.push(`<div class="ok">All trait names found as EBV_/ACC_ columns in ${esc(file)}.</div>`);
  }else if(file && SCAN.xlsx && SCAN.xlsx.includes(file)){
    parts.push(`<div class="err">${esc(file)} has no EBV_&lt;trait&gt;/ACC_&lt;trait&gt; column pairs.</div>`);
  }
  const html = parts.join('');
  if(box.dataset.html !== html){ box.innerHTML = html; box.dataset.html = html; }   // no DOM change when nothing changed
}

/* ---------- scan ---------- */
function fillList(id, arr){ $(id).innerHTML = (arr||[]).map(f => `<option value="${esc(f)}">`).join(''); }
function setIf(id, v){ if(v) $(id).value = v; }

async function scan(){
  const dir = $('s1_dir').value.trim();
  const m = $('scanmsg');
  if(!dir){ m.textContent = 'Type the folder path first.'; return; }
  m.textContent = 'Scanning…';
  try{
    const r = await api('/api/scan', {dir});
    if(!r.ok){ m.innerHTML = '<span style="color:var(--err)">'+esc(r.error)+'</span>'; return; }
    SCAN = r;
    fillList('dl_log', r.logs); fillList('dl_xlsx', r.xlsx); fillList('dl_dat', r.dats);
    fillList('dl_ped', r.peds); fillList('dl_si', r.xlsx);
    setIf('s1_log', r.suggest.log); setIf('s1_ebv', r.suggest.ebv);
    setIf('s1_dat', r.suggest.dat); setIf('s2_ped', r.suggest.ped);
    if(r.suggest.si) setIf('s2_si', r.suggest.si);
    $('s2_dir').value = dir;
    m.innerHTML = `Valid folder: <b>${esc(r.path)}</b> — ${r.files.length} files` +
      (r.dats.length > 1 ? `; ${r.dats.length} .dat files found, choose the right one in “Data file (.dat)”.` : '') +
      (r.dats.length === 0 ? '; <span style="color:var(--warn)">no .dat file found</span>.' : '.');
    showEffects();
    const cols = r.ebv_columns && r.ebv_columns[$('s1_ebv').value];
    if(cols && ENV.supports_traits) detectTraits();
    dimCheck(); updateOutPrev(); saveForm();
  }catch(e){ m.innerHTML = '<span style="color:var(--err)">'+esc(e.message)+'</span>'; }
}
function showEffects(){
  const log = $('s1_log').value.trim();
  const eff = (SCAN && SCAN.effects && SCAN.effects[log]) || [];
  const hint = SCAN ? SCAN.animal_hint : null;
  const box = $('efchips'); box.innerHTML = '';
  if(!SCAN){ return; }
  if(!eff.length){ $('efhint').innerHTML = `<span style="color:var(--warn)">No “Genetic variance(s) for effect N” block found in ${esc(log||'the log file')}.</span>`; return; }
  const ren = SCAN.renadd.map(x => x.file).join(', ');
  const dims = SCAN.dims[log] || {};
  $('efhint').innerHTML = `The log has genetic covariance matrices for effects: <b>${eff.map(e => e + (dims[e] ? ` (${dims[e]}×${dims[e]})` : '')).join(', ')}</b>.` +
    (hint ? ` Suggested animal effect: <b>${hint}</b> (from ${esc(ren)}).` : (ren ? ` Renumbered pedigree files: ${esc(ren)}.` : ' Choose the animal (additive) effect of your model.'));
  eff.forEach(n => {
    const b = document.createElement('button'); b.type = 'button'; b.className = 'ef';
    b.textContent = 'effect ' + n + (n === hint ? ' · animal' : '');
    b.onclick = () => { $('s1_effect').value = n; markEff(); dimCheck(); saveForm(); };
    b.dataset.n = n; box.appendChild(b);
  });
  // scanning is explicit: take the suggested animal effect, or clear a value that is not in this log
  if(hint) $('s1_effect').value = hint;
  else if(!eff.includes(parseInt($('s1_effect').value, 10))) $('s1_effect').value = '';
  markEff();
}
function markEff(){
  document.querySelectorAll('#efchips button').forEach(b => b.classList.toggle('sel', b.dataset.n === String($('s1_effect').value)));
}

/* ---------- output folder preview (step 2) ---------- */
function updateOutPrev(){
  const dir = $('s2_dir').value.trim().replace(/\/+$/,'');
  const name = $('s2_name').value.trim();
  const root = $('s2_root').value.trim() || ENV.workdir || '.';
  let out = name;
  if(!out){
    const seg = dir.split('/').filter(Boolean);
    if(seg.length >= 2){
      const leaf = seg[seg.length-1], parent = seg[seg.length-2];
      let tag = $('s2_tag').value.trim();
      if(!tag){ const p = leaf.split('_'); tag = (p.length >= 2 && p[0].toLowerCase() === 'blup') ? p[1] : leaf; }
      out = parent + '_' + tag + '_OCS';
    }
  }
  $('outprev').textContent = out ? 'Results will be saved in: ' + root.replace(/\/+$/,'') + '/' + out : '';
}

/* ---------- running ---------- */
function setStep(n, state){ $('st'+n).className = 'step ' + state; }
function setStatus(n, text, cls){ const s = $('stat'+n); s.textContent = text; s.className = 'status ' + (cls||''); }
function showMsg(n, text, kind){ const m = $('msg'+n); m.textContent = text || ''; m.className = 'msg ' + (text ? kind : ''); }

function params1(){
  syncFromDom();
  return {dir:$('s1_dir').value, effect:$('s1_effect').value, log:$('s1_log').value, ebv:$('s1_ebv').value,
    dat:$('s1_dat').value, out_xlsx:$('s1_outx').value, out_pdf:$('s1_outp').value,
    traits: TR.map(t => ({name: t.name.trim(), weight: t.w, direction: t.dir})),
    ref_min_year:$('s1_refyear').value, sd_mode:$('s1_sd').value, min_acc:$('s1_minacc').value,
    acc_traits:$('s1_acctraits').value, min_n:$('s1_minn').value, dat_id_col:$('s1_idcol').value,
    dat_year_col:$('s1_ycol').value};
}
function params2(){
  return {dir:$('s2_dir').value, ped:$('s2_ped').value, si:$('s2_si').value, sires:$('s2_sires').value,
    dams:$('s2_dams').value, birth_year:$('s2_year').value, cand:$('s2_cand').value,
    min_matings:$('s2_minm').value, max_matings:$('s2_maxm').value, ne:$('s2_ne').value,
    si_column:$('s2_sicol').value, alternatives:$('s2_alt').value, pop_size:$('s2_pop').value,
    sheet:$('s2_sheet').value, tag:$('s2_tag').value, out_root:$('s2_root').value, out_name:$('s2_name').value};
}

function lineEl(t){
  const d = document.createElement('div');
  if(/^ERROR|Error/.test(t)) d.className = 'e'; else if(/^WARNING|Warning/.test(t)) d.className = 'w'; else if(t.startsWith('$ ')) d.className = 'c';
  d.textContent = t; return d;
}

async function run(n){
  showMsg(n, '');
  const step = n === 1 ? 'si' : 'ocs';
  const log = $('log'+n); log.innerHTML = ''; log.style.display = 'block';
  $('res'+n).style.display = 'none'; $('res'+n).innerHTML = '';
  saveForm(); saveTraits();
  try{
    const r = await api('/api/run', {step, params: n === 1 ? params1() : params2()});
    JOB[n] = r.job; setStep(n, 'running'); setStatus(n, 'Running…');
    $('run'+n).disabled = true; $('stop'+n).disabled = false;
    poll(n, 0);
  }catch(e){
    log.style.display = 'none'; showMsg(n, e.message, 'err'); setStep(n, 'error'); setStatus(n, '');
  }
}

async function poll(n, off){
  try{
    const r = await api(`/api/job?id=${JOB[n]}&offset=${off}`);
    const log = $('log'+n);
    r.lines.forEach(t => log.appendChild(lineEl(t)));
    if(r.lines.length) log.scrollTop = log.scrollHeight;
    if(r.status === 'running'){ TIMER[n] = setTimeout(() => poll(n, r.next), 600); return; }
    $('run'+n).disabled = false; $('stop'+n).disabled = true;
    if(r.status === 'done'){
      setStep(n, 'done'); setStatus(n, `Done (${r.seconds}s)`, 'ok'); renderResult(n, r);
      if(n === 1) afterStep1();
    }else if(r.status === 'stopped'){
      setStep(n, ''); setStatus(n, 'Stopped');
    }else{
      setStep(n, 'error'); setStatus(n, `Failed (exit code ${r.rc})`, 'err');
      const errs = [...log.querySelectorAll('.e')].map(e => e.textContent);
      showMsg(n, errs.length ? errs[errs.length-1] : 'The job failed; see the log below.', 'err');
    }
  }catch(e){ showMsg(n, e.message, 'err'); $('run'+n).disabled = false; $('stop'+n).disabled = true; }
}

function renderTable(t){
  let h = `<h3>${esc(t.title)}${t.total ? ` <span class="status">(${t.rows.length}/${t.total})</span>` : ''}</h3><div class="tw"><table><thead><tr>` +
    t.columns.map(c => `<th>${esc(c)}</th>`).join('') + '</tr></thead><tbody>';
  t.rows.forEach(r => { h += '<tr>' + r.map(c => `<td>${esc(c)}</td>`).join('') + '</tr>'; });
  return h + '</tbody></table></div>';
}
function renderResult(n, r){
  let h = '';
  if(r.outputs.length){
    h += '<h3>Result files</h3><div class="files">' + r.outputs.map(o =>
      `<a href="/file?path=${encodeURIComponent(o.path)}" target="_blank">${esc(o.name)}</a>`).join('') + '</div>';
  }
  r.tables.forEach(t => h += renderTable(t));
  const box = $('res'+n); box.innerHTML = h; box.style.display = 'block';
}
function afterStep1(){
  setIf('s2_si', $('s1_outx').value);
  $('s2_dir').value = $('s1_dir').value;
  updateOutPrev(); saveForm();
  showMsg(1, 'Step 1 finished. The input folder and SI file were passed on to step 2; scroll down to run optiSel.', 'info');
}

async function stop(n){ if(JOB[n]) await api('/api/stop', {id: JOB[n]}); }

/* ---------- wiring ---------- */
$('scan').onclick = scan;
$('tadd').onclick = () => setTraitCount(Math.min(ENV.max_traits, TR.length + 1));
$('tdel').onclick = () => setTraitCount(Math.max(1, TR.length - 1));
$('tdetect').onclick = () => { const m = detectTraits(); $('wwarn').innerHTML = `<span>${esc(m)}</span>`; setTimeout(updateShares, 6000); };
$('wreset').onclick = resetWeights;
$('dimcheck').addEventListener('click', e => { const b = e.target.closest('button[data-fix]'); if(b) setTraitCount(+b.dataset.fix); });
$('s1_dir').addEventListener('keydown', e => { if(e.key === 'Enter') scan(); });
$('run1').onclick = () => run(1); $('run2').onclick = () => run(2);
$('stop1').onclick = () => stop(1); $('stop2').onclick = () => stop(2);
$('s1_log').addEventListener('change', () => { showEffects(); dimCheck(); });
$('s1_ebv').addEventListener('change', dimCheck);
$('s1_effect').addEventListener('input', () => { markEff(); dimCheck(); });
document.querySelectorAll('[data-save]').forEach(e => e.addEventListener('change', saveForm));
['s2_dir','s2_tag','s2_root','s2_name'].forEach(id => $(id).addEventListener('input', updateOutPrev));
$('theme').onclick = () => {
  const r = document.documentElement;
  const dark = r.getAttribute('data-theme') ? r.getAttribute('data-theme') === 'dark' : matchMedia('(prefers-color-scheme:dark)').matches;
  r.setAttribute('data-theme', dark ? 'light' : 'dark');
};
loadForm();
loadEnv().then(() => { loadTraits(); renderTraits(); updateOutPrev(); if($('s1_dir').value) scan(); });
</script>
</body>
</html>
"""


# ---------------------------------------------------------------- main
def main():
    global WORKDIR
    ap = argparse.ArgumentParser(description="Local web tool: BLUP -> selection index -> optiSel")
    ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--no-browser", action="store_true", help="do not open the browser automatically")
    ap.add_argument("--scripts-dir", default=str(Path(__file__).resolve().parent),
                    help="folder with si_cal.py and optiSel_pig.R (default: folder of this file)")
    ap.add_argument("--si-script", help="path of the SI script (default: si_cal.py or si_all.py)")
    ap.add_argument("--ocs-script", help="path of the optiSel script (default: optiSel_pig.R or ocs_pig.R)")
    ap.add_argument("--rscript", help="path of Rscript (default: found automatically)")
    ap.add_argument("--workdir", help="folder used to resolve relative paths (default: current folder)")
    a = ap.parse_args()

    if a.workdir:
        WORKDIR = Path(a.workdir).expanduser().resolve()
    sd = Path(a.scripts_dir).expanduser()
    CFG["si_script"] = str(Path(a.si_script).expanduser().resolve()) if a.si_script else \
        find_script(sd, ["si_cal.py", "si_all.py"])
    CFG["ocs_script"] = str(Path(a.ocs_script).expanduser().resolve()) if a.ocs_script else \
        find_script(sd, ["optiSel_pig.R", "ocs_pig.R"])
    CFG["rscript"] = find_rscript(a.rscript)
    if CFG["si_script"]:
        (CFG["default_traits"], CFG["default_directions"], CFG["max_traits"],
         CFG["supports_traits"], CFG["supports_weights"]) = read_script_config(CFG["si_script"])

    server = None
    for port in range(a.port, a.port + 20):
        try:
            server = ThreadingHTTPServer(("127.0.0.1", port), Handler)
            break
        except OSError:
            continue
    if server is None:
        sys.exit(f"Could not open a port in {a.port}-{a.port + 19}.")

    url = f"http://127.0.0.1:{server.server_address[1]}"
    print("=" * 60)
    print(" Pig breeding workflow web tool")
    print("=" * 60)
    print(f" Address         : {url}")
    print(f" Working folder  : {WORKDIR}")
    print(f" Step 1 (SI)     : {CFG['si_script'] or 'NOT FOUND'}")
    print(f" Step 2 (optiSel): {CFG['ocs_script'] or 'NOT FOUND'}")
    print(f" Rscript         : {CFG['rscript'] or 'NOT FOUND'}")
    print(f" Traits          : up to {CFG['max_traits']}"
          f"{'' if CFG['supports_traits'] else '  (si_cal.py is an old version: custom traits disabled)'}")
    print(" Press Ctrl+C to quit.")
    if not a.no_browser:
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nBye.")


if __name__ == "__main__":
    main()
